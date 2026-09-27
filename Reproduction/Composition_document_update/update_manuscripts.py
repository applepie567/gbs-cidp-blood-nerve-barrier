"""Integrate the composition feasibility result and the conditional external model."""
from pathlib import Path
from copy import deepcopy
import json, sys
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt
from docx.enum.text import WD_COLOR_INDEX

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT.parent/'validation_20260927/outputs'
O=ROOT/'outputs';O.mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT.parent/'validation_20260927/scripts'))
from update_manuscripts import borders as old_borders

def replace(p,text):
    prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for c in list(p._p):
        if c.tag!=qn('w:pPr'):p._p.remove(c)
    r=p.add_run(text)
    if prop is not None:r._r.insert(0,prop)
    r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    return p

def before(anchor,text,template=None,style=None):
    el=OxmlElement('w:p');anchor._p.addprevious(el);p=Paragraph(el,anchor._parent)
    if template is not None and template._p.pPr is not None:p._p.append(deepcopy(template._p.pPr))
    if style:p.style=style
    return replace(p,text)

S12CAP=("(A) A macrophage marker index predicts the observed macrophage nucleus fraction when one entire atlas donor is excluded from marker selection. "
"(B) The same procedure excludes the donor's whole centre. Both panels show 37 donors from GSE285983. Blue denotes Münster, teal Essen and orange Würzburg. "
"Markers were selected from donor balanced lineage profiles. The index uses within sample gene ranks and excludes the 11 Fc outcome genes. It measures relative marker expression and is not an estimated cell percentage. "
"(C) The fixed 20 gene index and Fc score in the independent GSE213455 biopsies. Blue denotes four CIDP cases and orange nine vasculitic neuropathy cases. Each point is one donor. "
"(D) CIDP minus vasculitic neuropathy Fc score differences with HC3 95% confidence intervals. Blue denotes the original unadjusted comparison, teal the model including the marker index, and grey the corresponding model using mean probe aggregation. "
"The conditional model has ten residual degrees of freedom. All intervals include zero. Supplementary Data 6 provides marker selections, held out predictions, donor scores and the plotted model estimates.")

def main():
    d=Document(SRC/'Manuscript_refocused.docx');p=list(d.paragraphs)
    replace(p[7],p[7].text.replace('An independent whole nerve comparison also had a lower CIDP score (g = −1.20, exact P = 0.076).','An independent whole nerve comparison had a lower CIDP score (g = −1.20, exact P = 0.076), and adjustment for a separate macrophage marker index reduced this difference.'))
    replace(p[15],p[15].text.replace('; Supplementary Data 5','. Supplementary Data 5')+' Supplementary Data 6 contains the nerve macrophage marker analysis and its source data.')
    before(p[38],"We then tested whether macrophage contribution could help explain the independent tissue result. A 20 gene marker index excluded the Fc genes and tracked macrophage nucleus fractions in held out atlas donors (Spearman ρ = 0.933). In GSE213455, the index correlated with the Fc score (ρ = 0.907). Including it in the disease model changed the Fc difference from −0.837 to 0.208 (95% confidence interval −0.916 to 1.332). Mean probe aggregation gave a similar conditional estimate. The tissue difference therefore depended on accounting for the broader macrophage expression profile (Supplementary Fig. S12 and Supplementary Table S6).",template=p[36])
    replace(p[41],p[41].text.replace('In nerve, the macrophage Fc receptor difference became smaller after composition adjustment.','In nerve, the Fc receptor difference became smaller after accounting for macrophage states in the atlas and for a macrophage marker index in the independent tissue cohort.'))
    replace(p[45],"The external nerve analysis broadens the comparison to vasculitic neuropathy and helps explain the tissue signal. Fc scores were lower in CIDP both in atlas macrophages and in the independent whole nerve cohort. In the external biopsies, Fc expression closely followed a macrophage marker index, and the conditional disease difference was small and imprecise. This pattern is consistent with differences in the macrophage contribution to whole tissue expression. The markers can also change with cell state, so the analysis does not isolate cell abundance. The external cohort contained only four CIDP cases, while centre imbalance affected the internal comparator check. A cohort with measured cell proportions and expression within macrophages would distinguish these explanations more directly.")
    replace(p[51],p[51].text.replace('In CIDP nerve, a lower Fc score also appeared in an independent tissue cohort, and the atlas analysis showed sensitivity to macrophage composition.','In nerve, the Fc differences depended on macrophage composition in the atlas and on a macrophage marker index in an independent tissue cohort.'))
    before(p[82],"A further exploratory analysis asked whether the external Fc difference depended on a broader macrophage expression profile. We derived markers from donor level counts in the nerve atlas and excluded all 11 Fc genes. Marker selection was repeated with each donor and each centre held out. After the index met the specified calibration criteria, its final 20 genes were fixed and scored by within sample ranks in GSE213455. We added this single index to the diagnosis model and used HC3 confidence intervals. Supplementary Methods gives the selection rules, feasibility criteria and probe aggregation sensitivity. The index was not converted to an absolute cell fraction.",template=p[80])
    replace(p[94],p[94].text.replace('These blood additions are not yet included in the archived v2.4.0 release.','Supplementary Data 6 contains donor and lineage counts for the nerve reference, marker selections, composition benchmarks and external conditional models. The blood additions and composition analysis are not yet included in the archived v2.4.0 release.'))
    replace(p[96],p[96].text.replace('The original three cohort blood synthesis is retained;','Supplementary Data 6 supplies the composition analysis plan, reference reconstruction, held out evaluation and figure code. The original three cohort blood synthesis is retained.'))
    before(p[160],"Supplementary Data 6. Nerve macrophage marker analysis, including donor and lineage aggregate counts, the fixed and held out marker sets, feasibility results, external conditional models and Supplementary Figure S12 source data.",template=p[159])
    for par in d.paragraphs:
        updated=par.text.replace('Tables S1–S5','Tables S1–S6').replace('Figures S1–S11','Figures S1–S12').replace('Data 1–5','Data 1–6')
        if updated!=par.text:replace(par,updated)
    d.save(O/'Manuscript_refocused.docx')

    s=Document(SRC/'Supplementary_Information_refocused.docx');q=list(s.paragraphs)
    replace(s.tables[0].rows[3].cells[4].paragraphs[0], 'The atlas Fc difference decreases after composition adjustment. The external tissue difference also decreases after adjustment for the macrophage marker index, with an imprecise conditional estimate.')
    replace(s.tables[1].rows[11].cells[3].paragraphs[0], 'Independent Fc score comparison and macrophage marker sensitivity [3].')
    replace(q[6],q[6].text+' Supplementary Data 6 contains the nerve macrophage marker analysis, donor and lineage counts, calibration results and conditional external models.')
    methods=[
      ('Macrophage marker index in whole nerve tissue','Heading 2'),
      ('We added this exploratory analysis after the external Fc result was known. The dated plan fixed the marker rules and feasibility criteria before constructing the reference or examining marker associations in GSE213455. We recovered all 37 GSE285983 CellBender matrices and retained the 365,708 nuclei listed in the authors’ metadata. Counts were summed by donor into 13 broad lineages. Macro1 and Macro2 formed the macrophage group. Related Schwann, perineurial, blood endothelial and mural clusters were pooled within their respective lineages. The remaining groups were endoneurial and epineurial stroma, lymphatic endothelium, T and NK cells, B cells, mast cells, granulocytes and adipocytes. Reference profiles required at least 20 nuclei for a donor and lineage. Whole donor mixtures retained every annotated nucleus. Reconstructed macrophage counts and Fc gene totals exactly matched the preceding analysis.','Normal'),
      ('Gene symbols were matched to unambiguous probes present on GPL13369. After excluding the 11 Fc genes, 16,633 common genes defined the ranking background. Mitochondrial and ribosomal genes, XIST and the Y linked markers listed in the script were ineligible as marker genes. Each eligible donor and lineage profile was normalised to counts per million. Lineage means gave equal weight to donors. Candidate markers required a macrophage mean of at least 10 CPM, expression above 1 CPM in at least half the macrophage donors, and at least fourfold enrichment over the largest other lineage mean after adding 1 CPM. The 20 most enriched genes were selected, with a minimum of ten required. The index was the mean within sample percentile rank of these genes among the common ranking background.','Normal'),
      ('We rebuilt the marker set after excluding each entire donor, then calculated the index in that donor’s mixture. The calibration targets were the observed macrophage nucleus fraction and the macrophage share of captured counts. The analysis plan required Spearman correlation of at least 0.60 with nucleus fraction and 0.70 with captured count share, with positive correlations within every centre containing at least five donors. We also excluded each whole centre and required positive correlations with both targets in all three centres. These were feasibility criteria for this extension, not established performance standards. External adjustment would be omitted if a criterion failed.','Normal'),
      ('After these criteria were met, the 20 markers derived from all atlas donors were fixed. Their percentile ranks were calculated separately in each of the 13 eligible external biopsies, using median probe expression. The existing Fc outcome and its standardisation were unchanged. The conditional model included an intercept, diagnosis and one standardised marker index. HC3 covariance and ten residual t degrees of freedom supplied its interval and model P value. Sensitivities used mean probe aggregation and omitted each donor while keeping the original score scale. No additional clinical covariates were added to this small model. The models address the same exploratory question and are reported together without selecting a favourable specification. Calibration in the atlas does not establish accuracy across assay platforms. Marker expression may reflect both cell abundance and cell state.','Normal')]
    for text,style in methods:before(q[35],text,template=q[33] if style=='Normal' else q[32],style=style)
    results=[
      ('Macrophage marker calibration and external conditional analysis','Heading 2'),
      ('The index met the specified feasibility criteria. All donor and centre exclusions retained 20 marker genes. Across the 37 held out donors, its Spearman correlations were 0.933 with macrophage nucleus fraction and 0.950 with captured count share. Within centre correlations with nucleus fraction ranged from 0.855 to 0.937. Excluding entire centres gave an overall correlation of 0.911 with nucleus fraction and 0.927 with captured count share. Every within centre correlation remained positive. Full marker lists, including CD163, MS4A7, CSF1R, F13A1, C1QA and C1QB, are supplied in Supplementary Data 6. None overlaps the 11 gene Fc outcome.','Normal'),
      ('In the 13 external biopsies, the fixed index correlated with the Fc score at ρ = 0.907. The unadjusted CIDP minus vasculitic neuropathy difference of −0.837 became 0.208 after including the index (HC3 95% confidence interval −0.916 to 1.332, model P = 0.689). Mean probe aggregation gave 0.151 (−0.752 to 1.053, P = 0.718). Donor omission coefficients ranged from 0.016 to 0.636 on the fixed score scale. The marker index ranged from 0.557 to 0.615 in CIDP and 0.574 to 0.682 in vasculitic neuropathy, with partial overlap. Maximum model leverage was 0.567. The small sample and correlated expression summaries leave the conditional difference imprecise. These results show that the tissue Fc difference depends on the broader macrophage expression profile. They do not distinguish abundance from state related expression (Supplementary Fig. S12 and Supplementary Table S6).','Normal')]
    for text,style in results:before(q[92],text,template=q[90] if style=='Normal' else q[89],style=style)
    replace(q[66],q[66].text.replace('These additions are supplied','Supplementary Data 6 contains the new nerve composition analysis and Figure S12. These additions are supplied'))
    before(q[73],"The composition analysis can be rerun from the donor and lineage count array, metadata, external probe matrix and annotations in Supplementary Data 6. The download script retrieves the 37 original matrices and full nucleus metadata if aggregation is repeated. Complete gene counts, all held out marker selections and panel source tables are retained. A separate base R implementation reproduced the conditional coefficient, HC3 interval and model P value.",template=q[71])
    # Add the compact three-line model table immediately before the figures.
    anchor=q[132]
    title=before(anchor,'Supplementary Table S6. External Fc models including a macrophage marker index',template=q[130],style='Heading 2')
    table=s.add_table(rows=1,cols=4)
    title._p.addnext(table._tbl)
    headers=['Model','Difference','95% confidence interval','Model P']
    for cell,txt in zip(table.rows[0].cells,headers):replace(cell.paragraphs[0],txt)
    data=[['Unadjusted','−0.837','−1.764 to 0.090','0.073'],['Macrophage index adjusted','0.208','−0.916 to 1.332','0.689'],['Mean probe sensitivity','0.151','−0.752 to 1.053','0.718']]
    for row in data:
        for cell,txt in zip(table.add_row().cells,row):replace(cell.paragraphs[0],txt)
    for j,width in enumerate([2.35,.8,2.05,.8]):
        table.columns[j].width=Inches(width)
        for row in table.rows:row.cells[j].width=Inches(width)
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_before=Pt(4);par.paragraph_format.space_after=Pt(4);par.paragraph_format.line_spacing=1
                for run in par.runs:run.font.size=Pt(10);run.bold=i==0
    old_borders(table)
    foot=before(anchor,'All models use the same four CIDP and nine vasculitic neuropathy biopsies. Differences are CIDP minus vasculitic neuropathy. Intervals and P values use HC3 covariance with 11 residual degrees of freedom for the unadjusted model and ten for the conditional models. The original exact rank P value is 0.0755 and is reported separately in Table S3. The mean probe sensitivity uses the marker index and Fc score obtained by averaging probes. The index is a gene expression summary, not an absolute cell fraction.',template=q[131],style='Normal')
    for run in foot.runs:run.bold=False;run.font.size=Pt(10)
    p12=s.add_paragraph('Supplementary Figure S12. Macrophage marker calibration and external Fc sensitivity',style='Heading 1')
    p12.paragraph_format.page_break_before=True
    for run in p12.runs:run.font.highlight_color=WD_COLOR_INDEX.YELLOW
    pic=s.add_paragraph();pic.add_run().add_picture(str(O/'Figures/Supplementary_Figure_12.png'),width=Inches(6.45))
    cap=s.add_paragraph(S12CAP,style='Normal');
    for run in cap.runs:run.font.highlight_color=WD_COLOR_INDEX.YELLOW;run.bold=False;run.font.size=Pt(10)
    s.save(O/'Supplementary_Information_refocused.docx')
    for name in ['Manuscript_refocused.docx','Supplementary_Information_refocused.docx']:
        doc=Document(O/name)
        (ROOT/'results'/name.replace('.docx','_text.txt')).write_text('\n'.join(p.text for p in doc.paragraphs))

if __name__=='__main__':main()
